"""Remove style tags that are absent from the English runtime strings."""

from pathlib import Path
import re
import sys

from openpyxl import load_workbook


DEFAULT_OFFSETS = {"442703", "443017", "443439", "447998", "1097765"}
STYLE_TAG_RE = re.compile(r"#[15]")


def main():
    path = Path(sys.argv[1])
    offsets = set(sys.argv[2:]) or DEFAULT_OFFSETS
    workbook = load_workbook(path)
    sheet = workbook.active
    headers = {cell.value: cell.column for cell in sheet[1]}
    offset_column = headers["offset"]
    korean_column = headers["ko"]
    changed = []

    for row in range(2, sheet.max_row + 1):
        offset = str(sheet.cell(row, offset_column).value or "")
        if offset not in offsets:
            continue
        cell = sheet.cell(row, korean_column)
        original = str(cell.value or "")
        updated = STYLE_TAG_RE.sub("", original)
        if updated == original:
            raise RuntimeError(f"{offset}: no removable style tag found")
        cell.value = updated
        changed.append(offset)

    if set(changed) != offsets:
        missing = sorted(offsets - set(changed))
        raise RuntimeError(f"offsets not updated: {', '.join(missing)}")

    workbook.save(path)
    print(f"updated {len(changed)} rows: {', '.join(changed)}")


if __name__ == "__main__":
    main()
