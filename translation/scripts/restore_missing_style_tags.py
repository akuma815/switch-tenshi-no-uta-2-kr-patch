"""Restore #5/#1 renderer tags from a missing-tag review CSV."""

import csv
from pathlib import Path
import re
import sys

from openpyxl import load_workbook


SPEAKER_TEMPLATE_RE = re.compile(r"^#5[^\\]+\\n#1")
SPEAKER_LINE_RE = re.compile(r"^(.+?:)\\n")


def restore_tags(english, korean):
    # The over-three-lines workbook can already contain a restored copy of a
    # master row. Make the operation idempotent and repair an interrupted
    # earlier run that doubled adjacent style markers.
    korean = korean.replace("#5#5", "#5").replace("#1#1", "#1")
    if [match.group(0) for match in re.finditer(r"#[15]", korean)] == ["#5", "#1"]:
        return korean

    if SPEAKER_TEMPLATE_RE.match(english):
        # Normal dialogue: color the translated speaker name, then reset the
        # body at the same first-line boundary as the English string.
        if SPEAKER_LINE_RE.match(korean):
            return "#5" + SPEAKER_LINE_RE.sub(r"\1\\n#1", korean, count=1)

        # A rare English line names a speaker although the Korean translation
        # contains only the body. Preserve the control sequence without
        # inventing a speaker name or adding a visible blank line.
        return "#5#1" + korean

    # Item and sign messages are yellow from start to end. Keep a trailing
    # renderer terminator (such as \\c) outside the #1 reset marker.
    suffix = ""
    if korean.endswith(r"\c"):
        korean, suffix = korean[:-2], r"\c"
    return "#5" + korean + "#1" + suffix


def main():
    workbook_path = Path(sys.argv[1])
    review_path = Path(sys.argv[2])
    with open(review_path, "r", encoding="utf-8-sig", newline="") as review:
        rows = {row["offset"]: row for row in csv.DictReader(review)}

    workbook = load_workbook(workbook_path)
    sheet = workbook.active
    headers = {cell.value: cell.column for cell in sheet[1]}
    offset_column = headers["offset"]
    english_column = headers["en"]
    korean_column = headers["ko"]
    changed = []

    for row_number in range(2, sheet.max_row + 1):
        offset = str(sheet.cell(row_number, offset_column).value or "")
        if offset not in rows:
            continue
        ko_cell = sheet.cell(row_number, korean_column)
        english = str(sheet.cell(row_number, english_column).value or "")
        korean = str(ko_cell.value or "")
        updated = restore_tags(english, korean)
        if updated != korean:
            ko_cell.value = updated
            changed.append(offset)

    workbook.save(workbook_path)
    print(f"updated {len(changed)} rows in {workbook_path.name}")


if __name__ == "__main__":
    main()
