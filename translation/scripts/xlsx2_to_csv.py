"""2편 번역 XLSX 4종 -> 빌드용 CSV

사용법: python xlsx2_to_csv.py <translation 폴더> <대사 출력 csv>
- 대사: 천사의시2_대사_번역.xlsx -> <대사 출력 csv>
- 크레딧/UI/게임선택: xlsx -> 백업_csv/*.csv 덮어쓰기 (빌드 스크립트가 여기서 읽음)
"""
import csv
import os
import sys

from openpyxl import load_workbook

base = sys.argv[1]
out_master_csv = sys.argv[2]
backup = sys.argv[3] if len(sys.argv) > 3 else os.path.join(base, "백업_csv")
os.makedirs(backup, exist_ok=True)


def read_xlsx(path, cols):
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    rows = []
    header = None
    for r in ws.iter_rows(values_only=True):
        if header is None:
            header = [str(c) if c is not None else "" for c in r]
            continue
        d = {}
        for i, name in enumerate(header):
            if name in cols:
                v = r[i] if i < len(r) else None
                d[name] = "" if v is None else str(v)
        rows.append(d)
    wb.close()
    return rows


def write_csv(path, cols, rows):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for r in rows:
            w.writerow({c: r.get(c, "") for c in cols})


# 1) 대사 마스터
COLS = ["offset", "key_hex", "jp", "en", "ko"]
# The runtime supplies this shop prompt's values in the English template
# order (%d price, then %s item name), unlike the Japanese source order.
# Keep the leading '*' control character while matching the call site.
RUNTIME_KO_OVERRIDES = {
    "28589": "*매입 가격은 %d G입니다.\\n%s 매입하시겠습니까?",
}
master = read_xlsx(os.path.join(base, "천사의시2_대사_번역.xlsx"), COLS)
print(f"master rows: {len(master)}")

for row in master:
    override = RUNTIME_KO_OVERRIDES.get(row.get("offset"))
    if override:
        row["ko"] = override
        print(f"runtime template normalized: offset {row['offset']}")

write_csv(out_master_csv, COLS, master)
print(f"written: {out_master_csv}")

# 2) 크레딧 / UI / 게임선택 -> 백업_csv
for xlsx, csv_name, cols in [
    ("천사의시2_크레딧_번역.xlsx", "visualscene2_paired.csv", ["offset", "jp", "en", "ko"]),
    ("천사의시2_UI_번역.xlsx", "ui_menu_text2_paired.csv", ["table", "index", "jp", "en", "ko"]),
    ("천사의시2_게임선택_번역.xlsx", "outgameinfo2_paired.csv", ["key", "offset", "en", "ko"]),
]:
    rows = read_xlsx(os.path.join(base, xlsx), cols)
    write_csv(os.path.join(backup, csv_name), cols, rows)
    ko_n = sum(1 for r in rows if (r.get("ko") or "").strip())
    print(f"{csv_name}: {len(rows)} rows (ko {ko_n})")
